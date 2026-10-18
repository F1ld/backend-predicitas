"""Pruebas de contratos/modulo-05.json (M05-E01 a M05-E10), F-09 y D-013. Casos CP-M05-01 a 09."""
import io
import json
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.main import app
from app.modules.auditoria.models import AuditoriaEvento
from app.modules.citas.models import Cita, HistorialEstadoCita
from app.modules.intervenciones import reglas
from app.modules.intervenciones.router import modo_sombra
from app.modules.lista_espera.models import CupoLiberado
from tests.conftest import cabecera
from tests.intervenciones.conftest import (HOY, NUEVA, crear_cita, crear_intervencion,
                                           crear_prediccion)


def _codigo(r):
    return r.json()["error"]["codigo"]


def _eventos(db, accion, modulo="INTERVENCION"):
    return list(db.scalars(select(AuditoriaEvento).where(AuditoriaEvento.accion == accion,
                                                         AuditoriaEvento.modulo == modulo)
                           .order_by(AuditoriaEvento.id_log)))


@pytest.fixture
def h(client, usuarios):
    return cabecera(client, usuarios["RECEPCION"])


# ---------------------------------------------------------------- reglas puras

def test_reglas_limite_y_respuesta():
    assert reglas.puede_crear_otra([]) and reglas.puede_crear_otra(["FALLIDA"])
    assert not reglas.puede_crear_otra(["FALLIDA", "FALLIDA"])
    assert not reglas.puede_crear_otra(["PENDIENTE"]) and not reglas.puede_crear_otra(["FALLIDA", "ENTREGADA"])
    assert reglas.admite_respuesta("ENTREGADA", "SMS") and reglas.admite_respuesta("ENVIADA", "LLAMADA_TELEFONICA")
    assert not reglas.admite_respuesta("ENVIADA", "SMS") and not reglas.admite_respuesta("PENDIENTE", "LLAMADA_TELEFONICA")
    assert reglas.canales_de("NO_EXISTE") == ()


# ---------------------------------------------------------------- M05-E01 elegibles (CP-M05-01)

def test_elegibles_filtra_y_ordena(client, h, db):
    medio = crear_cita(db, "MEDIO", id_cita_externa="CIT-TST-MEDIO", fecha_hora_cita="2026-10-26T09:00:00-05:00")
    alto = crear_cita(db, "ALTO", id_cita_externa="CIT-TST-ALTO", fecha_hora_cita="2026-10-28T09:00:00-05:00")
    reintento = crear_cita(db, "ALTO", id_cita_externa="CIT-TST-REIN", paciente_anon_id="PAC-TST-R")
    crear_intervencion(db, reintento, "FALLIDA")
    excluidas = [
        crear_cita(db, "BAJO", id_cita_externa="CIT-TST-BAJO", paciente_anon_id="PAC-TST-B"),
        crear_cita(db, None, id_cita_externa="CIT-TST-SINP", paciente_anon_id="PAC-TST-S"),
        crear_cita(db, "ALTO", id_cita_externa="CIT-TST-PAS", paciente_anon_id="PAC-TST-P",
                   fecha_hora_cita="2026-10-24T07:00:00-05:00"),  # ya pasó (hoy 08:00)
        crear_cita(db, "ALTO", id_cita_externa="CIT-TST-CANC", paciente_anon_id="PAC-TST-C", estado_actual="CANCELADA"),
        crear_cita(db, "ALTO", id_cita_externa="CIT-TST-LEJOS", paciente_anon_id="PAC-TST-L",
                   fecha_hora_cita="2026-11-05T09:00:00-05:00"),  # fuera de hoy + 7
    ]
    llegada = crear_cita(db, "ALTO", id_cita_externa="CIT-TST-LLEGO", paciente_anon_id="PAC-TST-LL")
    crear_intervencion(db, llegada, "ENTREGADA")
    # opt-out en otra cita del mismo paciente
    crear_cita(db, "ALTO", id_cita_externa="CIT-TST-OPT1", paciente_anon_id="PAC-TST-OPT")
    opt = db.query(Cita).filter_by(id_cita_externa="CIT-TST-OPT1").one()
    crear_intervencion(db, opt, "RESPONDIDA", "SIN_RESPUESTA", optout=True)
    crear_cita(db, "ALTO", id_cita_externa="CIT-TST-OPT2", paciente_anon_id="PAC-TST-OPT")
    # la predicción más reciente manda: ALTO antiguo, BAJO reciente → no elegible
    vieja = crear_cita(db, "ALTO", id_cita_externa="CIT-TST-BAJO2", paciente_anon_id="PAC-TST-B2")
    crear_prediccion(db, vieja, "BAJO", "0.100", cuando=HOY)

    r = client.get("/intervenciones/elegibles", headers=h)
    assert r.status_code == 200, r.text
    items = r.json()["items"]
    assert [i["id_cita_externa"] for i in items] == ["CIT-TST-REIN", "CIT-TST-ALTO", "CIT-TST-MEDIO"]
    assert items[0]["intervenciones_previas"] == 1 and items[0]["riesgo"]["resultado_simulado"] is True
    assert {"CIT-TST-I0001"} & {i["id_cita_externa"] for i in items} == set()
    assert len(excluidas) == 5 and medio and alto

    r = client.get("/intervenciones/elegibles", headers=h, params={"nivel_riesgo": "MEDIO"})
    assert [i["id_cita_externa"] for i in r.json()["items"]] == ["CIT-TST-MEDIO"]
    r = client.get("/intervenciones/elegibles", headers=h,
                   params={"desde": "2026-10-24", "hasta": "2026-11-06", "especialidad": "CARDIOLOGIA", "limite": 10})
    assert "CIT-TST-LEJOS" in {i["id_cita_externa"] for i in r.json()["items"]}
    r = client.get("/intervenciones/elegibles", headers=h, params={"nivel_riesgo": "BAJO"})
    assert r.status_code == 422


def test_elegibles_solo_escritores(client, usuarios):
    r = client.get("/intervenciones/elegibles", headers=cabecera(client, usuarios["TECNICO_IA"]))
    assert r.status_code == 403 and _codigo(r) == "SIN_PERMISO"


# ---------------------------------------------------------------- M05-E02 crear (CP-M05-02, 05, 06, 08)

def test_crear_exige_consentimiento_y_audita(client, h, db):
    cita = crear_cita(db)
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita,
                                                         "consentimiento_confirmado": False})
    assert r.status_code == 422 and _codigo(r) == "CONSENTIMIENTO_REQUERIDO"
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita})
    assert r.status_code == 201, r.text
    d = r.json()
    assert d["estado_intervencion"] == "PENDIENTE" and d["id_cita_externa"] == cita.id_cita_externa
    assert d["acciones_disponibles"] == ["REGISTRAR_ENVIO", "REGISTRAR_FALLO"]
    evento = _eventos(db, "CREATE")[-1]
    assert evento.datos_after["consentimiento_confirmado"] is True and evento.datos_after["id_prediccion"]
    assert evento.entidad_afectada == "intervenciones" and evento.id_entidad == d["id_intervencion"]


@pytest.mark.parametrize(("cuerpo", "campo"), [
    ({"plantilla_mensaje": "LLAMADA_GUION_V1"}, "plantilla_mensaje"),  # no es de WhatsApp
    ({"plantilla_mensaje": "INVENTADA"}, "plantilla_mensaje"),
    ({"canal_contacto": "PALOMA"}, "canal_contacto"),
    ({"consentimiento_confirmado": "talvez"}, "consentimiento_confirmado"),
])
def test_crear_validaciones(client, h, db, cuerpo, campo):
    cita = crear_cita(db)
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita, **cuerpo})
    assert r.status_code == 422 and _codigo(r) == "VALIDACION" and campo in r.json()["error"]["campos"]


def test_crear_404_y_403(client, usuarios, h):
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": 999999999})
    assert r.status_code == 404 and _codigo(r) == "NO_ENCONTRADO"
    r = client.post("/intervenciones", headers=cabecera(client, usuarios["TECNICO_IA"]), json={**NUEVA, "id_cita": 1})
    assert r.status_code == 403


@pytest.mark.parametrize("cambios", [
    {"nivel": "BAJO"}, {"nivel": None}, {"estado_actual": "CANCELADA"},
    {"fecha_hora_cita": "2026-10-24T07:30:00-05:00"},  # pasada
])
def test_crear_no_elegible(client, h, db, cambios):
    nivel = cambios.pop("nivel", "ALTO")
    cita = crear_cita(db, nivel, **cambios)
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita})
    assert r.status_code == 409 and _codigo(r) == "NO_ELEGIBLE"


def test_optout_bloquea_todas_las_citas_del_paciente(client, h, db):
    primera = crear_cita(db, paciente_anon_id="PAC-TST-777")
    segunda = crear_cita(db, id_cita_externa="CIT-TST-I0002", paciente_anon_id="PAC-TST-777")
    inter = crear_intervencion(db, primera, "ENTREGADA")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/optout", headers=h)
    assert r.status_code == 200 and r.json()["consentimiento_optout"] is True
    assert len(_eventos(db, "OPTOUT")) == 1
    client.post(f"/intervenciones/{inter.id_intervencion}/optout", headers=h)  # idempotente
    assert len(_eventos(db, "OPTOUT")) == 1
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": segunda.id_cita})
    assert r.status_code == 409 and _codigo(r) == "OPTOUT_VIGENTE"


def test_limite_una_que_llega_y_un_reintento(client, h, db):
    cita = crear_cita(db)
    cuerpo = {**NUEVA, "id_cita": cita.id_cita}
    primera = client.post("/intervenciones", headers=h, json=cuerpo).json()["id_intervencion"]
    r = client.post("/intervenciones", headers=h, json=cuerpo)
    assert r.status_code == 409 and _codigo(r) == "LIMITE_INTERVENCION"  # hay una en curso
    client.post(f"/intervenciones/{primera}/estado", headers=h, json={"estado_intervencion": "FALLIDA"})
    r = client.post("/intervenciones", headers=h, json=cuerpo)
    assert r.status_code == 201  # reintento único tras FALLIDA
    segunda = r.json()["id_intervencion"]
    for estado in ("ENVIADA", "ENTREGADA"):
        client.post(f"/intervenciones/{segunda}/estado", headers=h, json={"estado_intervencion": estado})
    r = client.post("/intervenciones", headers=h, json=cuerpo)
    assert r.status_code == 409 and _codigo(r) == "LIMITE_INTERVENCION"


def test_dos_fallidas_agotan_el_reintento(client, h, db):
    cita = crear_cita(db)
    crear_intervencion(db, cita, "FALLIDA")
    crear_intervencion(db, cita, "FALLIDA")
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita})
    assert r.status_code == 409 and _codigo(r) == "LIMITE_INTERVENCION"


def test_modo_sombra_bloquea_la_creacion(client, h, db):
    cita = crear_cita(db)
    app.dependency_overrides[modo_sombra] = lambda: True
    r = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita})
    assert r.status_code == 409 and _codigo(r) == "MODO_SOMBRA"


def test_modo_sombra_lee_la_configuracion(monkeypatch):
    from app.core.config import get_settings
    from app.modules.intervenciones import router
    monkeypatch.setattr(get_settings(), "intervencion_modo_sombra", True)
    assert router.modo_sombra() is True


# ---------------------------------------------------------------- M05-E03 / E04 lectura

def test_listar_y_detalle(client, usuarios, db):
    cita = crear_cita(db)
    otra = crear_cita(db, id_cita_externa="CIT-TST-I0002", paciente_anon_id="PAC-TST-X")
    vieja = crear_intervencion(db, cita, "FALLIDA", fecha_hora_envio=HOY - timedelta(days=2))
    nueva = crear_intervencion(db, cita, "RESPONDIDA", "CONFIRMA", canal="SMS", fecha_hora_envio=HOY - timedelta(days=1))
    crear_intervencion(db, otra, "PENDIENTE", fecha_hora_envio=HOY)
    h = cabecera(client, usuarios["TECNICO_IA"])

    r = client.get("/intervenciones", headers=h, params={"id_cita": cita.id_cita})
    assert [i["id_intervencion"] for i in r.json()["items"]] == [nueva.id_intervencion, vieja.id_intervencion]
    assert client.get("/intervenciones", headers=h).json()["total"] == 3
    for params, total in (({"estado_intervencion": "FALLIDA"}, 1), ({"canal_contacto": "SMS"}, 1),
                          ({"respuesta_paciente": "CONFIRMA"}, 1), ({"desde": "2026-10-24"}, 1),
                          ({"hasta": "2026-10-22"}, 1)):
        assert client.get("/intervenciones", headers=h, params=params).json()["total"] == total, params

    d = client.get(f"/intervenciones/{nueva.id_intervencion}", headers=h).json()
    assert d["acciones_disponibles"] == ["REGISTRAR_OPTOUT", "CONFIRMAR_CITA"]
    r = client.get("/intervenciones/999999999", headers=h)
    assert r.status_code == 404 and _codigo(r) == "NO_ENCONTRADO"


@pytest.mark.parametrize(("estado", "canal", "respuesta", "optout", "cita", "esperado"), [
    ("ENVIADA", "SMS", None, False, "RECIBIDA", ["REGISTRAR_ENTREGA", "REGISTRAR_FALLO"]),
    ("ENVIADA", "LLAMADA_TELEFONICA", None, False, "RECIBIDA",
     ["REGISTRAR_ENTREGA", "REGISTRAR_FALLO", "REGISTRAR_RESPUESTA"]),
    ("ENTREGADA", "SMS", None, False, "RECIBIDA", ["REGISTRAR_RESPUESTA"]),
    ("RESPONDIDA", "SMS", "SOLICITA_LIBERACION", True, "CONFIRMADA", ["LIBERAR_CUPO"]),
    ("RESPONDIDA", "SMS", "CONFIRMA", False, "CONFIRMADA", ["REGISTRAR_OPTOUT"]),
    ("RESPONDIDA", "SMS", "SOLICITA_LIBERACION", True, "LIBERADA", []),
    ("FALLIDA", "SMS", None, False, "RECIBIDA", []),
])
def test_acciones_disponibles(estado, canal, respuesta, optout, cita, esperado):
    assert reglas.acciones_disponibles(estado, canal, respuesta, optout, cita) == esperado


# ---------------------------------------------------------------- M05-E05 envío (CP-M05-03)

def test_ciclo_de_envio_y_auditoria(client, h, db):
    cita = crear_cita(db)
    inter = crear_intervencion(db, cita)
    r = client.post(f"/intervenciones/{inter.id_intervencion}/estado", headers=h, json={"estado_intervencion": "ENVIADA"})
    assert r.status_code == 200 and r.json()["acciones_disponibles"] == ["REGISTRAR_ENTREGA", "REGISTRAR_FALLO"]
    assert r.json()["fecha_hora_envio"].startswith("2026-10-24T08:00")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/estado", headers=h, json={"estado_intervencion": "ENTREGADA"})
    assert r.status_code == 200 and r.json()["estado_intervencion"] == "ENTREGADA"
    updates = _eventos(db, "UPDATE")
    assert [(e.datos_before["estado_intervencion"], e.datos_after["estado_intervencion"]) for e in updates] == \
        [("PENDIENTE", "ENVIADA"), ("ENVIADA", "ENTREGADA")]
    r = client.post(f"/intervenciones/{inter.id_intervencion}/estado", headers=h, json={"estado_intervencion": "FALLIDA"})
    assert r.status_code == 409 and _codigo(r) == "TRANSICION_INVALIDA"


def test_envio_errores(client, h, db):
    inter = crear_intervencion(db, crear_cita(db))
    r = client.post(f"/intervenciones/{inter.id_intervencion}/estado", headers=h, json={"estado_intervencion": "ENTREGADA"})
    assert r.status_code == 409 and _codigo(r) == "TRANSICION_INVALIDA"
    r = client.post(f"/intervenciones/{inter.id_intervencion}/estado", headers=h, json={"estado_intervencion": "RESPONDIDA"})
    assert r.status_code == 422
    r = client.post("/intervenciones/999999999/estado", headers=h, json={"estado_intervencion": "ENVIADA"})
    assert r.status_code == 404


# ---------------------------------------------------------------- M05-E06 respuesta (CP-M05-04)

def test_respuesta_no_cambia_la_cita(client, h, db):
    cita = crear_cita(db)
    inter = crear_intervencion(db, cita, "ENTREGADA")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/respuesta", headers=h,
                    json={"respuesta_paciente": "CONFIRMA"})
    assert r.status_code == 200, r.text
    assert r.json()["estado_intervencion"] == "RESPONDIDA" and r.json()["fecha_hora_respuesta"]
    assert client.get(f"/citas/{cita.id_cita}", headers=h).json()["estado_actual"] == "RECIBIDA"
    historial = db.scalars(select(HistorialEstadoCita).where(HistorialEstadoCita.id_cita == cita.id_cita)).all()
    assert [x.estado_nuevo for x in historial] == ["RECIBIDA"]
    assert len(_eventos(db, "RESPUESTA")) == 1 and _eventos(db, "OPTOUT") == []


def test_respuesta_con_optout_y_por_llamada(client, h, db):
    inter = crear_intervencion(db, crear_cita(db), "ENVIADA", canal="LLAMADA_TELEFONICA")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/respuesta", headers=h,
                    json={"respuesta_paciente": "SIN_RESPUESTA", "optout": True})
    assert r.status_code == 200 and r.json()["consentimiento_optout"] is True
    assert len(_eventos(db, "RESPUESTA")) == 1 and len(_eventos(db, "OPTOUT")) == 1


def test_respuesta_transicion_invalida(client, h, db):
    inter = crear_intervencion(db, crear_cita(db), "ENVIADA", canal="SMS")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/respuesta", headers=h,
                    json={"respuesta_paciente": "CONFIRMA"})
    assert r.status_code == 409 and _codigo(r) == "TRANSICION_INVALIDA"
    assert client.post("/intervenciones/999999999/respuesta", headers=h,
                       json={"respuesta_paciente": "CONFIRMA"}).status_code == 404
    assert client.post("/intervenciones/999999999/optout", headers=h).status_code == 404


# ---------------------------------------------------------------- M05-E08 liberar (CP-M05-07)

def test_liberar_cupo_atomico(client, h, db):
    cita = crear_cita(db)
    inter = crear_intervencion(db, cita, "RESPONDIDA", "SOLICITA_LIBERACION")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/liberar-cupo", headers=h,
                    json={"motivo": "Paciente cede su cupo voluntariamente"})
    assert r.status_code == 201, r.text
    c = r.json()
    assert (c["estado_cupo"], c["origen_liberacion"], c["medico"], c["especialidad"]) == \
        ("DISPONIBLE", "INTERVENCION_VOLUNTARIA", "Dr. Prueba", "CARDIOLOGIA")
    assert c["id_cita_externa_origen"] == cita.id_cita_externa and c["oferta_vigente"] is None
    db.refresh(cita)
    assert cita.estado_actual == "LIBERADA"
    hist = db.scalars(select(HistorialEstadoCita).where(HistorialEstadoCita.id_cita == cita.id_cita)
                      .order_by(HistorialEstadoCita.id_historial)).all()
    assert hist[-1].estado_nuevo == "LIBERADA" and hist[-1].motivo_cambio.startswith("Paciente cede")
    liberar = _eventos(db, "LIBERAR")
    assert len(liberar) == 1 and liberar[0].id_entidad == c["id_cupo_liberado"]

    r = client.post(f"/intervenciones/{inter.id_intervencion}/liberar-cupo", headers=h)
    assert r.status_code == 409 and _codigo(r) == "CUPO_YA_LIBERADO"


def test_liberar_cupo_errores(client, h, db):
    sin = crear_intervencion(db, crear_cita(db), "RESPONDIDA", "CONFIRMA")
    r = client.post(f"/intervenciones/{sin.id_intervencion}/liberar-cupo", headers=h)
    assert r.status_code == 409 and _codigo(r) == "SIN_SOLICITUD_LIBERACION"

    pasada = crear_cita(db, id_cita_externa="CIT-TST-PAS", fecha_hora_cita="2026-10-24T07:00:00-05:00")
    inter = crear_intervencion(db, pasada, "RESPONDIDA", "SOLICITA_LIBERACION")
    r = client.post(f"/intervenciones/{inter.id_intervencion}/liberar-cupo", headers=h)
    assert r.status_code == 409 and _codigo(r) == "CITA_PASADA"

    cancelada = crear_cita(db, id_cita_externa="CIT-TST-CAN", estado_actual="CANCELADA")
    inter = crear_intervencion(db, cancelada, "RESPONDIDA", "SOLICITA_LIBERACION")
    db.commit()
    r = client.post(f"/intervenciones/{inter.id_intervencion}/liberar-cupo", headers=h)
    assert r.status_code == 409 and _codigo(r) == "TRANSICION_INVALIDA"
    assert client.post("/intervenciones/999999999/liberar-cupo", headers=h).status_code == 404


def test_liberar_cupo_si_falla_el_cupo_la_cita_no_queda_liberada(client, h, db, monkeypatch):
    """RNF-014: si falla la creación del cupo, se revierte también la transición de la cita."""
    cita = crear_cita(db)
    inter = crear_intervencion(db, cita, "RESPONDIDA", "SOLICITA_LIBERACION")
    db.commit()  # punto de partida confirmado dentro de la transacción de prueba

    from app.modules.intervenciones import service

    def cupo_roto(**kwargs):
        raise RuntimeError("fallo simulado al crear el cupo")
    monkeypatch.setattr(service, "CupoLiberado", cupo_roto)
    try:
        r = client.post(f"/intervenciones/{inter.id_intervencion}/liberar-cupo", headers=h)
        assert r.status_code == 500 and _codigo(r) == "ERROR_INTERNO"
    except RuntimeError:  # TestClient puede propagar la excepción del servidor
        pass
    db.expire_all()
    assert db.get(Cita, cita.id_cita).estado_actual == "RECIBIDA"
    assert db.scalar(select(CupoLiberado).where(CupoLiberado.id_cita_origen == cita.id_cita)) is None
    hist = db.scalars(select(HistorialEstadoCita).where(HistorialEstadoCita.id_cita == cita.id_cita)).all()
    assert [x.estado_nuevo for x in hist] == ["RECIBIDA"]


# ---------------------------------------------------------------- M05-E09 plantillas

def test_plantillas(client, usuarios, h):
    r = client.get("/intervenciones/plantillas", headers=h)
    assert r.status_code == 200
    ids = {p["id"] for p in r.json()}
    assert ids == {"RECORDATORIO_CONFIRMACION_V1", "RECORDATORIO_LIBERACION_V1", "LLAMADA_GUION_V1"}
    assert all(p["contenido"] == "[PENDIENTE: contenido aprobado]" for p in r.json())
    assert client.get("/intervenciones/plantillas", headers=cabecera(client, usuarios["TECNICO_IA"])).status_code == 403


# ---------------------------------------------------------------- M05-E10 REP-02 (CP-M05-09)

def _datos_rep02(db):
    cita = crear_cita(db)
    otra = crear_cita(db, "MEDIO", id_cita_externa="CIT-TST-I0002", paciente_anon_id="PAC-TST-X",
                      especialidad="PEDIATRIA")
    crear_intervencion(db, cita, "FALLIDA", fecha_hora_envio=HOY - timedelta(hours=5))
    crear_intervencion(db, cita, "RESPONDIDA", "SOLICITA_LIBERACION", fecha_hora_envio=HOY - timedelta(hours=3),
                       fecha_hora_respuesta=HOY - timedelta(hours=2, minutes=30))
    crear_intervencion(db, otra, "RESPONDIDA", "CONFIRMA", canal="SMS", fecha_hora_envio=HOY - timedelta(hours=2),
                       fecha_hora_respuesta=HOY - timedelta(hours=1))
    crear_intervencion(db, otra, "ENTREGADA", canal="SMS", fecha_hora_envio=HOY - timedelta(hours=1))
    crear_intervencion(db, otra, "PENDIENTE", canal="SMS", fecha_hora_envio=HOY)
    db.add(CupoLiberado(id_cita_origen=cita.id_cita, especialidad="CARDIOLOGIA", medico="Dr. Prueba",
                        fecha_hora_cupo=cita.fecha_hora_cita, estado_cupo="DISPONIBLE",
                        origen_liberacion="INTERVENCION_VOLUNTARIA"))
    db.flush()


def test_rep02_json(client, h, db):
    _datos_rep02(db)
    r = client.get("/reportes/rep-02", headers=h, params={"desde": "2026-10-01", "hasta": "2026-10-31"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["resumen"] == {"enviadas": 4, "entregadas": 3, "respondidas": 2, "confirmaciones": 1,
                            "liberaciones": 1, "tasa_respuesta": 0.667}
    assert d["resultado_simulado"] is True and len(d["filas"]) == 5
    assert {"canal_contacto": "SMS", "respuesta_paciente": "CONFIRMA", "cantidad": 1} in d["por_canal"]
    tiempos = {f["respuesta_paciente"]: f["tiempo_respuesta_minutos"] for f in d["filas"] if f["respuesta_paciente"]}
    assert tiempos == {"SOLICITA_LIBERACION": 30, "CONFIRMA": 60}
    assert [f["cupo_liberado"] for f in d["filas"]].count(True) == 1
    for params, n in (({"canal_contacto": "SMS"}, 3), ({"nivel_riesgo": "MEDIO"}, 3),
                      ({"especialidad": "CARDIOLOGIA"}, 2), ({"hasta": "2026-10-01"}, 0)):
        assert len(client.get("/reportes/rep-02", headers=h, params=params).json()["filas"]) == n, params
    assert _eventos(db, "EXPORT", "REPORTES") == []  # json no se audita


def test_rep02_csv_y_xlsx_auditan_export(client, h, db):
    _datos_rep02(db)
    r = client.get("/reportes/rep-02", headers=h, params={"formato": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    lineas = r.content.decode("utf-8-sig").splitlines()
    assert lineas[0] == ("id_intervencion,id_cita_externa,especialidad,nivel_riesgo,canal_contacto,"
                         "estado_intervencion,respuesta_paciente,tiempo_respuesta_minutos,cupo_liberado")
    assert len(lineas) == 6
    r = client.get("/reportes/rep-02", headers=h, params={"formato": "xlsx", "especialidad": "PEDIATRIA"})
    hoja = load_workbook(io.BytesIO(r.content)).active
    assert hoja.max_row == 4 and hoja.cell(1, 1).value == "id_intervencion"
    eventos = _eventos(db, "EXPORT", "REPORTES")
    assert [e.datos_after["formato"] for e in eventos] == ["csv", "xlsx"]
    assert eventos[1].datos_after["filtros"]["especialidad"] == "PEDIATRIA"
    assert "paciente" not in json.dumps([e.datos_after for e in eventos])


def test_rep02_errores(client, usuarios, h):
    r = client.get("/reportes/rep-02", headers=h, params={"formato": "pdf"})
    assert r.status_code == 422 and _codigo(r) == "FORMATO_NO_DISPONIBLE"
    r = client.get("/reportes/rep-02", headers=h, params={"canal_contacto": "FAX"})
    assert r.status_code == 422 and _codigo(r) == "VALIDACION"
    r = client.get("/reportes/rep-02", headers=cabecera(client, usuarios["TECNICO_IA"]))
    assert r.status_code == 403
    r = client.get("/reportes/rep-02", headers=h)
    assert r.json()["resumen"]["tasa_respuesta"] == 0.0


# ---------------------------------------------------------------- privacidad de la auditoría

def test_auditoria_sin_datos_personales(client, h, db):
    cita = crear_cita(db)
    i = client.post("/intervenciones", headers=h, json={**NUEVA, "id_cita": cita.id_cita}).json()["id_intervencion"]
    for estado in ("ENVIADA", "ENTREGADA"):
        client.post(f"/intervenciones/{i}/estado", headers=h, json={"estado_intervencion": estado})
    client.post(f"/intervenciones/{i}/respuesta", headers=h, json={"respuesta_paciente": "SOLICITA_LIBERACION",
                                                                     "optout": True})
    client.post(f"/intervenciones/{i}/liberar-cupo", headers=h)
    texto = json.dumps([[e.datos_before, e.datos_after] for e in db.scalars(
        select(AuditoriaEvento).where(AuditoriaEvento.modulo == "INTERVENCION"))])
    for prohibido in ("nombre", "correo", "telefono", "password", "token", "Dr. Prueba"):
        assert prohibido not in texto
