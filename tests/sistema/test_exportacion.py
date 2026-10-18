"""app.core.exportacion.responder: JSON, CSV con BOM y XLSX (contratos/comun.json → exportacion, D-016)."""
import io
import json
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.errores import Validacion
from app.core.exportacion import responder

COLUMNAS = ["id_cita", "especialidad", "probabilidad_no_show", "fecha_hora_cita", "medico"]
FILAS = [
    {"id_cita": 1, "especialidad": "CARDIOLOGIA", "probabilidad_no_show": Decimal("0.512"),
     "fecha_hora_cita": datetime(2026, 10, 24, 14, 0, tzinfo=timezone.utc), "medico": "Dr. Núñez",
     "no_pedida": "x"},
    {"id_cita": 2, "especialidad": "PEDIATRIA", "probabilidad_no_show": None,
     "fecha_hora_cita": None, "medico": "=HYPERLINK(1)"},
]


def test_json_devuelve_solo_las_columnas_pedidas():
    r = responder(FILAS, COLUMNAS, "json", "reporte")
    assert r.media_type == "application/json"
    datos = json.loads(r.body)
    assert datos[0] == {"id_cita": 1, "especialidad": "CARDIOLOGIA", "probabilidad_no_show": 0.512,
                        "fecha_hora_cita": "2026-10-24T14:00:00+00:00", "medico": "Dr. Núñez"}
    assert len(datos) == 2


def test_csv_utf8_con_bom_y_cabecera_del_contrato():
    r = responder(FILAS, COLUMNAS, "csv", "citas_por_riesgo")
    assert r.body.startswith(b"\xef\xbb\xbf")
    assert r.headers["content-type"].startswith("text/csv")
    disposicion = r.headers["content-disposition"]
    assert disposicion.startswith('attachment; filename="citas_por_riesgo_') and disposicion.endswith('.csv"')
    lineas = r.body.decode("utf-8-sig").splitlines()
    assert lineas[0] == ",".join(COLUMNAS)
    assert lineas[1] == "1,CARDIOLOGIA,0.512,2026-10-24T14:00:00+00:00,Dr. Núñez"
    assert lineas[2] == "2,PEDIATRIA,,,'=HYPERLINK(1)"  # se neutraliza la fórmula


def test_xlsx_valido_con_cabecera_en_negrita():
    r = responder(FILAS, COLUMNAS, "XLSX", "reporte")
    assert r.headers["content-disposition"].endswith('.xlsx"')
    hoja = load_workbook(io.BytesIO(r.body)).active
    assert [c.value for c in hoja[1]] == COLUMNAS
    assert all(c.font.bold for c in hoja[1])
    assert hoja["A2"].value == 1 and hoja["C2"].value == 0.512
    assert hoja["D2"].value == datetime(2026, 10, 24, 9, 0)  # hora de Lima
    assert hoja["E3"].value == "'=HYPERLINK(1)"
    assert hoja.max_row == 3


def test_formato_desconocido_es_validacion():
    with pytest.raises(Validacion) as exc:
        responder(FILAS, COLUMNAS, "pdf", "reporte")
    assert exc.value.campos == {"formato": "Debe ser json, csv o xlsx"}
